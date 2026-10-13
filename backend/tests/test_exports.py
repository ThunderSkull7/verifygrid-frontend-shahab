"""Automated test suite for verification results retrieval and report generation."""

import io
import uuid
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.core.config import settings


@pytest.fixture
def upload_storage(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    root = tmp_path / "uploads"
    monkeypatch.setattr(settings, "storage_root", root)
    return root


def create_project(client: TestClient, name: str = "Exports Test Project") -> str:
    response = client.post("/api/v1/projects", json={"name": name})
    assert response.status_code == 201
    return response.json()["id"]


def upload_file(client: TestClient, project_id: str, filename: str, content: bytes) -> dict:
    response = client.post(
        f"/api/v1/projects/{project_id}/files",
        files={"file": (filename, content, "text/csv")},
    )
    assert response.status_code == 201
    return response.json()


def create_test_run(client: TestClient, project_id: str) -> str:
    csv_1 = (
        "id, name, amount\n"
        "101, Alice, 1000\n"
        "102, Bob, 2000\n"
        "103, Charlie, 3000\n"
    ).encode("utf-8")
    csv_2 = (
        "id, name, amount\n"
        "101, Alice, 1000\n"
        "102, Bob, 2500\n"   # mismatch
        "104, Dave, 4000\n"  # missing in source
    ).encode("utf-8")

    f1 = upload_file(client, project_id, "src.csv", csv_1)
    f2 = upload_file(client, project_id, "tgt.csv", csv_2)

    res = client.post(
        f"/api/v1/projects/{project_id}/runs",
        json={
            "source_file_id": f1["id"],
            "target_file_id": f2["id"],
            "key_columns": ["id"],
        },
    )
    assert res.status_code == 201
    return res.json()["id"]


def test_get_run_summary_endpoint(client: TestClient, upload_storage: Path):
    project_id = create_project(client, "Summary Test Project")
    run_id = create_test_run(client, project_id)

    res = client.get(f"/api/v1/projects/{project_id}/runs/{run_id}/summary")
    assert res.status_code == 200
    summary = res.json()
    assert summary["matched_records"] == 2
    assert summary["identical_matches"] == 1
    assert summary["field_mismatches"] == 1
    assert summary["unmatched_source_records"] == 1


def test_export_csv_report(client: TestClient, upload_storage: Path):
    project_id = create_project(client, "CSV Export Project")
    run_id = create_test_run(client, project_id)

    res = client.get(f"/api/v1/projects/{project_id}/runs/{run_id}/export?format=csv")
    assert res.status_code == 200
    assert "text/csv" in res.headers["content-type"]
    assert f"attachment; filename=\"verification_report_{run_id}.csv\"" in res.headers["content-disposition"]

    content = res.content.decode("utf-8")
    assert "# VerifyGrid Verification Report" in content
    assert f"# Run ID,{run_id}" in content
    assert "Category,Match Status,Source Row,Target Row,Key Values,Field Differences,Message" in content
    assert "field_mismatch" in content
    assert "missing_record" in content


def test_export_xlsx_report(client: TestClient, upload_storage: Path):
    project_id = create_project(client, "XLSX Export Project")
    run_id = create_test_run(client, project_id)

    res = client.get(f"/api/v1/projects/{project_id}/runs/{run_id}/export?format=xlsx")
    assert res.status_code == 200
    assert "openxmlformats-officedocument.spreadsheetml.sheet" in res.headers["content-type"]
    assert f"attachment; filename=\"verification_report_{run_id}.xlsx\"" in res.headers["content-disposition"]

    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    assert "Summary" in wb.sheetnames
    assert "Discrepancies" in wb.sheetnames

    ws_summary = wb["Summary"]
    assert ws_summary["A1"].value == "VerifyGrid Reconciliation Summary Report"

    ws_disc = wb["Discrepancies"]
    assert ws_disc.cell(row=1, column=1).value == "Category"
    assert ws_disc.max_row >= 4  # Header + 3 discrepancies


def test_export_empty_discrepancies_report(client: TestClient, upload_storage: Path):
    project_id = create_project(client, "Clean Project")
    csv_bytes = "id,val\n1,A\n".encode("utf-8")
    f1 = upload_file(client, project_id, "clean.csv", csv_bytes)

    run_res = client.post(
        f"/api/v1/projects/{project_id}/runs",
        json={
            "source_file_id": f1["id"],
            "target_file_id": f1["id"],
            "key_columns": ["id"],
        },
    )
    assert run_res.status_code == 201
    run_id = run_res.json()["id"]

    # Export CSV with 0 discrepancies
    res_csv = client.get(f"/api/v1/projects/{project_id}/runs/{run_id}/export?format=csv")
    assert res_csv.status_code == 200
    assert "# Total Discrepancies,0" in res_csv.content.decode("utf-8")

    # Export XLSX with 0 discrepancies
    res_xlsx = client.get(f"/api/v1/projects/{project_id}/runs/{run_id}/export?format=xlsx")
    assert res_xlsx.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(res_xlsx.content))
    assert "Summary" in wb.sheetnames
    assert "Discrepancies" in wb.sheetnames


def test_export_filtered_report(client: TestClient, upload_storage: Path):
    project_id = create_project(client, "Filtered Export Project")
    run_id = create_test_run(client, project_id)

    # Filter by discrepancy_category='field_mismatch'
    res = client.get(
        f"/api/v1/projects/{project_id}/runs/{run_id}/export?format=csv&discrepancy_category=field_mismatch"
    )
    assert res.status_code == 200
    content = res.content.decode("utf-8")
    assert "field_mismatch" in content
    assert "missing_record" not in content


def test_export_invalid_format_returns_422(client: TestClient, upload_storage: Path):
    project_id = create_project(client, "Invalid Format Project")
    run_id = create_test_run(client, project_id)

    res = client.get(f"/api/v1/projects/{project_id}/runs/{run_id}/export?format=pdf")
    assert res.status_code == 422
    assert "Unsupported export format" in res.json()["detail"]


def test_export_cross_project_isolation(client: TestClient, upload_storage: Path):
    proj_a = create_project(client, "Isolation A")
    proj_b = create_project(client, "Isolation B")
    run_a = create_test_run(client, proj_a)

    # Attempting to export run_a from project_b returns 404
    res = client.get(f"/api/v1/projects/{proj_b}/runs/{run_a}/export?format=csv")
    assert res.status_code == 404
    assert "Verification run not found" in res.json()["detail"]


def test_discrepancies_sorting_and_validation(client: TestClient, upload_storage: Path):
    project_id = create_project(client, "Sorting Project")
    run_id = create_test_run(client, project_id)

    # Sort descending by source_row_number
    res = client.get(
        f"/api/v1/projects/{project_id}/runs/{run_id}/results?sort_by=source_row_number&sort_order=desc"
    )
    assert res.status_code == 200
    items = res.json()["items"]
    assert len(items) == 3

    # Invalid sort_by returns 422
    invalid_res = client.get(
        f"/api/v1/projects/{project_id}/runs/{run_id}/results?sort_by=invalid_field"
    )
    assert invalid_res.status_code == 422
    assert "Invalid sort_by" in invalid_res.json()["detail"]


def test_export_json_report(client: TestClient, upload_storage: Path):
    project_id = create_project(client, "JSON Export Project")
    run_id = create_test_run(client, project_id)

    res = client.get(f"/api/v1/projects/{project_id}/runs/{run_id}/export?format=json")
    assert res.status_code == 200
    assert "application/json" in res.headers["content-type"]
    assert f"attachment; filename=\"verification_report_{run_id}.json\"" in res.headers["content-disposition"]

    data = res.json()
    assert data["report_type"] == "VerifyGrid Verification Report"
    assert data["run"]["id"] == run_id
    assert data["run"]["status"] == "completed"
    assert data["summary"]["matched_records"] == 2
    assert data["discrepancies_count"] == 3
    assert len(data["discrepancies"]) == 3
    # Check fields of discrepancy items
    item = data["discrepancies"][0]
    assert "id" in item
    assert "match_status" in item
    assert "discrepancy_category" in item
    assert "key_values" in item


def test_export_json_empty_discrepancies(client: TestClient, upload_storage: Path):
    project_id = create_project(client, "JSON Clean Project")
    csv_bytes = "id,val\n1,A\n".encode("utf-8")
    f1 = upload_file(client, project_id, "clean.csv", csv_bytes)

    run_res = client.post(
        f"/api/v1/projects/{project_id}/runs",
        json={
            "source_file_id": f1["id"],
            "target_file_id": f1["id"],
            "key_columns": ["id"],
        },
    )
    assert run_res.status_code == 201
    run_id = run_res.json()["id"]

    res = client.get(f"/api/v1/projects/{project_id}/runs/{run_id}/export?format=json")
    assert res.status_code == 200
    data = res.json()
    assert data["discrepancies_count"] == 0
    assert data["discrepancies"] == []


def test_export_missing_run_returns_404(client: TestClient, upload_storage: Path):
    project_id = create_project(client, "Missing Run Project")
    fake_run_id = str(uuid.uuid4())

    res = client.get(f"/api/v1/projects/{project_id}/runs/{fake_run_id}/export?format=csv")
    assert res.status_code == 404
    assert "Verification run not found" in res.json()["detail"]


def test_export_incomplete_run_returns_400(client: TestClient, upload_storage: Path):
    project_id = create_project(client, "Incomplete Run Project")
    csv_bytes = "id,name\n1,Alice\n".encode("utf-8")
    f1 = upload_file(client, project_id, "f1.csv", csv_bytes)

    # Creating a run with an invalid key column results in status="failed"
    fail_res = client.post(
        f"/api/v1/projects/{project_id}/runs",
        json={
            "source_file_id": f1["id"],
            "target_file_id": f1["id"],
            "key_columns": ["invalid_col"],
        },
    )
    assert fail_res.status_code == 422

    # Query the runs list to find the failed run's ID
    list_res = client.get(f"/api/v1/projects/{project_id}/runs")
    assert list_res.status_code == 200
    runs = list_res.json()["items"]
    assert len(runs) == 1
    failed_run_id = runs[0]["id"]
    assert runs[0]["status"] == "failed"

    # Attempt export of the failed run -> 400 Bad Request
    export_res = client.get(f"/api/v1/projects/{project_id}/runs/{failed_run_id}/export?format=csv")
    assert export_res.status_code == 400
    assert "Cannot export verification run in 'failed' state" in export_res.json()["detail"]

